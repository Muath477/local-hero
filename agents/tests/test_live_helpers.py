"""The live tools are only useful if their offline parts are sound: the sample
document must exercise every code path Tarjuman claims to handle."""
import io

import docx
from docx.oxml.ns import qn

from benchmarks import smoke_e2e, tarjuman_live
from tools.tarjuman import translate_docx_with_stats


def test_sample_document_covers_every_structure():
    d = docx.Document(io.BytesIO(tarjuman_live.build_sample()))
    assert d.tables and d.sections[0].header.paragraphs[0].text and d.sections[0].footer.paragraphs[0].text
    assert any(p._p.findall(qn("w:hyperlink")) for p in d.paragraphs)
    assert any(r.bold for p in d.paragraphs for r in p.runs) and any(r.italic for p in d.paragraphs for r in p.runs)
    texts = [p.text for p in d.paragraphs]
    assert texts.count("Confidential") == 2, "repeated boilerplate exercises deduplication"


def test_sample_round_trips_and_translates_with_a_stub():
    src = tarjuman_live.build_sample()
    _, stats = translate_docx_with_stats(src, "French", translate=lambda t: t)
    assert stats == {}
    out, _ = translate_docx_with_stats(src, "Arabic", translate=lambda t: ["\u0646\u0635" if any(c.isalpha() for c in x) else x for x in t])
    doc = docx.Document(io.BytesIO(out))
    assert doc.tables[0].cell(0, 0).text == "\u0646\u0635"
    assert doc.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    assert len(doc.paragraphs[-2]._p.findall(qn("w:hyperlink"))) == 1


def test_smoke_cases_reference_real_skills_and_tools():
    from orchestrator.skills import load_skills
    from tools.registry import get_tool

    names = {s.name for s in load_skills()}
    for _, role, tool, skill, _text in smoke_e2e.CASES:
        assert skill is None or skill in names
        assert tool is None or get_tool(tool) is not None
