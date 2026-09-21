import io
import re

import docx
import pytest
import requests

import tools.builtin_tools as bt
import tools.file_reader as fr
from orchestrator import ollama_client
from tools import registry, tarjuman
from tools.builtin_tools import calculator, get_datetime, hf_model_search, list_uploaded_files
from tools.registry import search_tools


class TestRegistryUpgrades:
    @pytest.fixture(autouse=True)
    def clean(self, monkeypatch):
        monkeypatch.setattr(registry, "_REGISTRY", {})

    def test_schema_types_come_from_annotations(self):
        @registry.tool("demo", params={"n": "how many"})
        def demo(n: int, ratio: float, flag: bool, items: list[str], name: str = "x", raw=None):
            return n

        props = registry.all_schemas()[0]["function"]["parameters"]["properties"]
        assert [props[k]["type"] for k in ("n", "ratio", "flag", "items", "name", "raw")] == \
            ["integer", "number", "boolean", "array", "string", "string"]
        assert props["n"]["description"] == "how many"
        assert registry.all_schemas()[0]["function"]["parameters"]["required"] == ["n", "ratio", "flag", "items"]

    def test_string_arguments_are_coerced_to_the_declared_types(self):
        seen = {}

        @registry.tool("demo")
        def demo(n: int, ratio: float, flag: bool, name: str):
            seen.update(n=n, ratio=ratio, flag=flag, name=name)

        registry.get_tool("demo")(n="5", ratio="2.5", flag="true", name="7")
        assert seen == {"n": 5, "ratio": 2.5, "flag": True, "name": "7"}
        registry.get_tool("demo")(n="5.0", ratio=1, flag="no", name="x")
        assert seen["n"] == 5 and seen["flag"] is False

    def test_uncoercible_values_pass_through_for_the_function_to_reject(self):
        @registry.tool("demo")
        def demo(n: int):
            return n

        assert registry.get_tool("demo")(n="abc") == "abc"

    def test_arguments_the_function_does_not_declare_are_dropped(self):
        @registry.tool("demo")
        def demo(a: str):
            return a

        assert registry.get_tool("demo")(a="x", hallucinated="y") == "x"

    def test_var_kwargs_functions_receive_everything(self):
        @registry.tool("demo")
        def demo(**kw):
            return kw

        assert registry.get_tool("demo")(a=1, b=2) == {"a": 1, "b": 2}

    def test_decorated_function_is_still_directly_callable(self):
        @registry.tool("demo")
        def demo(n: int):
            return n + 1

        assert demo(1) == 2

    def test_arabic_keywords_find_english_described_tools(self):
        @registry.tool("Evaluate arithmetic", keywords=("احسب",))
        def calc(x: str):
            return x

        @registry.tool("Read a file")
        def other(x: str):
            return x

        assert search_tools("احسب لي ناتج 5+5", top_k=1)[0]["function"]["name"] == "calc"

    def test_only_and_exclude_scope_the_pool(self):
        for name in ("a_tool", "b_tool", "c_tool"):
            registry._REGISTRY[name] = {"fn": None, "keywords": [], "schema": {"function": {"name": name, "description": "x"}}}
        assert [s["function"]["name"] for s in search_tools("q", only=["b_tool"])] == ["b_tool"]
        assert "a_tool" not in [s["function"]["name"] for s in search_tools("q", exclude=["a_tool"])]


class TestRealToolSelection:
    """With the real registry: does an Arabic/English request surface the right tool in the top 5?"""

    @pytest.mark.parametrize("message, expected", [
        ("ابحث لي عن سعر الذهب", "web_search"),
        ("كم يساوي 15 × 23 + 7؟", "calculator"),
        ("what time is it in London right now?", "get_datetime"),
        ("وش الملفات اللي رفعتها؟", "list_uploaded_files"),
        ("ترجم لي هذي الجملة للعربية: hello", "translate_text"),
        ("ترجم الملف contract.docx للعربية", "translate_document"),
        ("find a GGUF model for Arabic on Hugging Face", "hf_model_search"),
        ("اقرأ الملف report.txt", "read_uploaded_file"),
    ])
    def test_expected_tool_is_offered(self, message, expected):
        from tools.agent_tools import DELEGATE_TOOLS
        names = [s["function"]["name"] for s in search_tools(message, exclude=DELEGATE_TOOLS)]
        assert expected in names and len(names) <= 5

    def test_delegation_tools_belong_to_the_council_only(self):
        from tools.agent_tools import DELEGATE_TOOLS
        assert {s["function"]["name"] for s in search_tools("x", top_k=len(DELEGATE_TOOLS), only=DELEGATE_TOOLS)} == set(DELEGATE_TOOLS)
        assert not set(DELEGATE_TOOLS) & {s["function"]["name"] for s in search_tools("x", exclude=DELEGATE_TOOLS)}


class TestCalculator:
    @pytest.mark.parametrize("expr, expected", [
        ("2+2", "4"), ("12.5 * (3 + 4)", "87.5"), ("2 ** 10", "1024"), ("7 // 2", "3"), ("7 % 4", "3"),
        ("sqrt(16)", "4"), ("round(3.14159, 2)", "3.14"), ("min(3, 1, 2) + max(4, 9)", "10"),
        ("15 × 23 + 7", "352"), ("١٢ + ٣", "15"), ("2 ^ 3", "8"), ("10 ÷ 4", "2.5"), ("-5 + +3", "-2"),
        ("pi * 2", "6.28318530718"), ("1/3", "0.333333333333"),
    ])
    def test_evaluates(self, expr, expected):
        assert calculator(expr) == expected

    @pytest.mark.parametrize("expr", [
        "__import__('os').system('dir')", "open('x')", "().__class__", "abs.__call__(1)", "x + 1", "'a' * 3",
        "[1, 2][0]", "lambda: 1", "2 ** 100000", "9 ** 9 ** 9", "sqrt(1, key=2)", "True + 1", "", "1 +",
    ])
    def test_refuses_anything_but_arithmetic(self, expr):
        assert calculator(expr).startswith("Error")

    def test_division_by_zero_is_a_message_not_a_crash(self):
        assert calculator("1/0") == "Error: division by zero."

    def test_overlong_expression(self):
        assert calculator("1+" * 200 + "1").startswith("Error")


class TestDatetime:
    def test_default_timezone_and_shape(self):
        out = get_datetime()
        assert re.match(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2} \(\w+ / .+\), Asia/Riyadh, UTC\+03:00$", out)

    def test_other_timezone(self):
        assert get_datetime("UTC").endswith("UTC, UTC+00:00")

    def test_unknown_timezone_is_reported(self):
        assert get_datetime("Mars/Olympus").startswith("Error: unknown timezone")


@pytest.fixture
def uploads(tmp_path, monkeypatch):
    d = tmp_path / "uploads"
    d.mkdir()
    monkeypatch.setattr(fr, "UPLOADS_DIR", d)
    return d


class TestUploadedFiles:
    def test_empty(self, uploads):
        assert list_uploaded_files() == "No files uploaded."

    def test_lists_visible_files_with_sizes(self, uploads):
        (uploads / "a.txt").write_text("hello")
        (uploads / ".gitkeep").write_text("")
        (uploads / "sub").mkdir()
        assert list_uploaded_files() == "- a.txt (5 bytes)"

    def test_missing_directory(self, tmp_path, monkeypatch):
        monkeypatch.setattr(fr, "UPLOADS_DIR", tmp_path / "nope")
        assert list_uploaded_files() == "No files uploaded."


def fake_arabic(monkeypatch, calls=None):
    def chat(model, messages, **kw):
        if calls is not None:
            calls.append(messages[-1]["content"])
        return {"message": {"content": "نص مترجم"}}

    monkeypatch.setattr(ollama_client, "chat", chat)


class TestTranslationTools:
    def test_translate_text_uses_the_engine_and_caches(self, monkeypatch, tmp_path):
        monkeypatch.setattr(tarjuman, "CACHE_PATH", tmp_path / "c.json")
        calls = []
        fake_arabic(monkeypatch, calls)
        assert bt.translate_text("The meeting starts at noon.", "Arabic") == "نص مترجم"
        assert bt.translate_text("The meeting starts at noon.", "Arabic") == "نص مترجم"
        assert len(calls) == 1, "the second identical request must come from the translation cache"

    def test_translate_text_rejects_huge_input(self):
        assert bt.translate_text("x" * (bt.MAX_TEXT_CHARS + 1)).startswith("Error")

    def test_translate_document_writes_a_new_file_and_keeps_the_original(self, uploads, monkeypatch, tmp_path):
        monkeypatch.setattr(tarjuman, "CACHE_PATH", tmp_path / "c.json")
        fake_arabic(monkeypatch)
        d = docx.Document()
        d.add_paragraph("Hello world")
        d.save(uploads / "memo.docx")
        out = bt.translate_document("memo.docx", "Arabic")
        assert "memo.docx" in out and "memo.arabic.docx" in out
        assert docx.Document(uploads / "memo.arabic.docx").paragraphs[0].text == "نص مترجم"
        assert docx.Document(uploads / "memo.docx").paragraphs[0].text == "Hello world"

    @pytest.mark.parametrize("name, fragment", [
        ("../secret.docx", "outside"), ("ghost.docx", "not found"), ("a.txt", "only .docx"),
    ])
    def test_translate_document_guards(self, uploads, tmp_path, name, fragment):
        (tmp_path / "secret.docx").write_bytes(b"x")
        (uploads / "a.txt").write_text("x")
        out = bt.translate_document(name, "Arabic")
        assert out.startswith("Error") and fragment in out


class TestHfSearch:
    class Resp:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            pass

        def json(self):
            return self.payload

    def test_formats_results(self, monkeypatch):
        seen = {}

        def get(url, params=None, timeout=None):
            seen.update(url=url, params=params)
            return self.Resp([{"id": "Qwen/Qwen3-4B-GGUF", "downloads": 426000, "likes": 161, "pipeline_tag": "text-generation"},
                              {"id": "x/y"}])

        monkeypatch.setattr(bt.requests, "get", get)
        out = hf_model_search("qwen3 gguf")
        assert "- Qwen/Qwen3-4B-GGUF — 426,000 downloads, 161 likes, text-generation" in out and "- x/y — 0 downloads" in out
        assert seen["params"]["search"] == "qwen3 gguf" and seen["params"]["sort"] == "downloads"

    def test_no_results_and_network_errors_are_strings(self, monkeypatch):
        monkeypatch.setattr(bt.requests, "get", lambda *a, **k: self.Resp([]))
        assert hf_model_search("zzz") == "No models found."

        def boom(*a, **k):
            raise requests.ConnectionError("offline")

        monkeypatch.setattr(bt.requests, "get", boom)
        assert hf_model_search("x").startswith("Hugging Face search failed")
