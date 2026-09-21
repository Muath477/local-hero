import docx
import pytest
import requests

import tools.file_reader as fr
import tools.web_search as ws
from tools import registry
from tools.mcp_bridge import _coerce_args


class TestRegistry:
    def test_builtin_tools_are_registered_with_valid_schemas(self):
        for name in ("web_search", "read_uploaded_file", "ask_writer_agent", "ask_coder_agent", "ask_researcher_agent"):
            assert registry.get_tool(name) is not None
        for schema in registry.all_schemas():
            fn = schema["function"]
            assert fn["name"] and fn["description"]
            params = fn["parameters"]
            assert params["type"] == "object"
            assert set(params.get("required", [])) <= set(params["properties"])

    def test_required_only_lists_params_without_defaults(self, monkeypatch):
        monkeypatch.setattr(registry, "_REGISTRY", {})

        @registry.tool("demo")
        def demo(a, b="x"):
            return a

        params = registry.all_schemas()[0]["function"]["parameters"]
        assert params["required"] == ["a"] and set(params["properties"]) == {"a", "b"}

    def test_search_tools_ranks_by_description_and_respects_top_k(self):
        hits = registry.search_tools("snippets", top_k=2)
        assert len(hits) <= 2
        assert hits[0]["function"]["name"] == "web_search"

    def test_unknown_tool_is_none(self):
        assert registry.get_tool("does_not_exist") is None


class TestFileReader:
    @pytest.fixture
    def uploads(self, tmp_path, monkeypatch):
        d = tmp_path / "uploads"
        d.mkdir()
        monkeypatch.setattr(fr, "UPLOADS_DIR", d)
        return d

    def test_reads_text(self, uploads):
        (uploads / "a.txt").write_text("مرحبا hello", encoding="utf-8")
        assert fr.read_uploaded_file("a.txt") == "مرحبا hello"

    def test_reads_docx(self, uploads):
        d = docx.Document()
        d.add_paragraph("first")
        d.add_paragraph("second")
        d.save(uploads / "a.docx")
        assert fr.read_uploaded_file("a.docx") == "first\nsecond"

    def test_missing_file(self, uploads):
        assert fr.read_uploaded_file("ghost.txt").startswith("Error:")

    def test_binary_file_is_reported_not_dumped(self, uploads):
        (uploads / "img.jpg").write_bytes(b"\xff\xd8\xff\xe0" + bytes(range(128, 256)))
        assert "binary" in fr.read_uploaded_file("img.jpg")

    @pytest.mark.parametrize("evil", ["../secret.txt", "..\\secret.txt", "../uploads_evil/x.txt"])
    def test_path_traversal_is_refused(self, uploads, tmp_path, evil):
        (tmp_path / "secret.txt").write_text("top secret")
        (tmp_path / "uploads_evil").mkdir()
        (tmp_path / "uploads_evil" / "x.txt").write_text("sibling dir with a shared prefix")
        out = fr.read_uploaded_file(evil)
        assert out.startswith("Error:") and "secret" not in out and "sibling" not in out

    def test_absolute_path_is_refused(self, uploads, tmp_path):
        outside = tmp_path / "outside.txt"
        outside.write_text("nope")
        assert fr.read_uploaded_file(str(outside)).startswith("Error:")


class TestCoerceArgs:
    SCHEMA = {
        "properties": {
            "engines": {"type": "array"},
            "maybe": {"anyOf": [{"type": "array"}, {"type": "null"}]},
            "query": {"type": "string"},
        }
    }

    @pytest.mark.parametrize(
        "value, expected",
        [("a,b", ["a", "b"]), ("a+b", ["a", "b"]), ("a b", ["a", "b"]), ("a, b, c", ["a", "b", "c"])],
    )
    def test_string_becomes_array(self, value, expected):
        assert _coerce_args(self.SCHEMA, {"engines": value})["engines"] == expected

    def test_anyof_array_is_recognised(self):
        assert _coerce_args(self.SCHEMA, {"maybe": "a+b"})["maybe"] == ["a", "b"]

    def test_leaves_well_typed_and_unrelated_args_alone(self):
        args = {"engines": ["x"], "query": "a, b", "unknown": "p+q"}
        assert _coerce_args(self.SCHEMA, args) == args

    def test_single_token_stays_a_string(self):
        assert _coerce_args(self.SCHEMA, {"engines": "duckduckgo"})["engines"] == "duckduckgo"


class TestWebSearch:
    HTML = """
    <div class="result"><a class="result__title">Python 3.14</a>
      <a class="result__url">python.org</a><div class="result__snippet">Released</div></div>
    <div class="result"><a class="result__title">Second</a></div>
    """
    LITE = """
    <table><tr><td><a class="result-link" href="https://python.org">Python 3.14 Lite</a></td></tr>
    <tr><td class="result-snippet">Latest release</td></tr>
    <tr><td><a class="result-link" href="https://b.example">Other</a></td></tr></table>
    """

    class _Resp:
        def __init__(self, text="", status=200):
            self.text, self.status_code = text, status

        def raise_for_status(self):
            pass

    def _sessions(self, monkeypatch, behaviour):
        """Fake requests.Session whose post(url, ...) is answered by behaviour(url); records the urls hit."""
        urls = []

        class S:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def post(self, url, *a, **k):
                urls.append(url)
                return behaviour(url)

        monkeypatch.setattr(ws.requests, "Session", S)
        monkeypatch.setattr(ws.time, "sleep", lambda s: None)
        return urls

    def test_parses_html_results(self, monkeypatch):
        urls = self._sessions(monkeypatch, lambda u: self._Resp(self.HTML))
        out = ws.web_search("python")
        assert "Python 3.14" in out and "python.org" in out and "Released" in out and "Second" in out
        assert urls == [ws.HTML_URL]

    def test_a_202_challenge_from_the_html_endpoint_falls_through_to_lite(self, monkeypatch):
        urls = self._sessions(monkeypatch, lambda u: self._Resp("", 202) if u == ws.HTML_URL else self._Resp(self.LITE))
        out = ws.web_search("latest version of python")
        assert "Python 3.14 Lite" in out and "https://python.org" in out and "Latest release" in out
        assert urls == [ws.HTML_URL, ws.LITE_URL], "an empty answer must move on, not repeat the same request"

    def test_an_empty_200_also_moves_on(self, monkeypatch):
        self._sessions(monkeypatch, lambda u: self._Resp("<html></html>") if u == ws.HTML_URL else self._Resp(self.LITE))
        assert "Lite" in ws.web_search("x")

    def test_network_error_on_one_endpoint_still_reaches_the_other(self, monkeypatch):
        def behaviour(u):
            if u == ws.HTML_URL:
                raise requests.ConnectionError("reset")
            return self._Resp(self.LITE)

        urls = self._sessions(monkeypatch, behaviour)
        assert "Lite" in ws.web_search("x") and urls.count(ws.HTML_URL) == ws.MAX_RETRIES

    def test_mcp_search_is_the_last_resort(self, monkeypatch):
        self._sessions(monkeypatch, lambda u: self._Resp("", 202))
        monkeypatch.setitem(ws._REGISTRY, "mcp_search_search", {
            "fn": lambda **kw: f"mcp result for {kw['q']}",
            "schema": {"function": {"parameters": {"required": ["q"]}}}})
        assert ws.web_search("python") == "mcp result for python"

    def test_a_broken_mcp_fallback_is_ignored(self, monkeypatch):
        self._sessions(monkeypatch, lambda u: self._Resp("", 202))

        def boom(**kw):
            raise RuntimeError("server died")

        monkeypatch.setitem(ws._REGISTRY, "mcp_search_search", {"fn": boom, "schema": {"function": {"parameters": {}}}})
        assert ws.web_search("x").startswith("No results found")

    def test_everything_empty_says_it_may_be_rate_limiting(self, monkeypatch):
        self._sessions(monkeypatch, lambda u: self._Resp("<html></html>"))
        monkeypatch.delitem(ws._REGISTRY, "mcp_search_search", raising=False)
        out = ws.web_search("zzz")
        assert out.startswith("No results found") and "rate limit" in out

    def test_network_failure_returns_string_after_retries_on_every_endpoint(self, monkeypatch):
        def boom(u):
            raise requests.ConnectionError("reset")

        urls = self._sessions(monkeypatch, boom)
        monkeypatch.delitem(ws._REGISTRY, "mcp_search_search", raising=False)
        out = ws.web_search("x")
        assert "failed after" in out and len(urls) == 2 * ws.MAX_RETRIES
